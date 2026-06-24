"""Connector base class, canonical transaction contract and mapping engine.

The default :meth:`BaseConnector.parse` is driven by a YAML mapping file so that
adjusting to a real exchange export is mostly a matter of editing column names
and the type-translation table — no code change. Subclasses override
:meth:`normalize_row` only for quirks that a flat mapping cannot express.
"""
from __future__ import annotations

from abc import ABC
from dataclasses import dataclass, field
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from typing import Any, BinaryIO

import yaml
from openpyxl import load_workbook

from app.core.money import to_decimal
from app.models.enums import AccountPlatform, TransactionType

MAPPINGS_DIR = Path(__file__).resolve().parent / "mappings"


@dataclass
class CanonicalTransaction:
    """Platform-agnostic, calculation-ready transaction.

    Convention: ``*_in`` is the asset received (acquisition side), ``*_out`` the
    asset given (disposal side). Symbols are upper-cased; amounts are Decimal.
    """

    external_id: str | None
    timestamp: datetime
    type: TransactionType
    asset_in: str | None = None
    amount_in: Decimal | None = None
    asset_out: str | None = None
    amount_out: Decimal | None = None
    fee_asset: str | None = None
    fee_amount: Decimal | None = None
    eur_value: Decimal | None = None
    price_source: str | None = None
    notes: str | None = None
    raw: dict[str, Any] = field(default_factory=dict)


@dataclass
class RowError:
    row: int
    message: str
    raw: dict[str, Any] = field(default_factory=dict)


@dataclass
class ParseResult:
    transactions: list[CanonicalTransaction] = field(default_factory=list)
    errors: list[RowError] = field(default_factory=list)


def load_mapping(name: str) -> dict[str, Any]:
    with open(MAPPINGS_DIR / name, encoding="utf-8") as fh:
        return yaml.safe_load(fh)


class BaseConnector(ABC):
    source: AccountPlatform
    name: str
    mapping_file: str

    def __init__(self) -> None:
        self.mapping = load_mapping(self.mapping_file)

    # --- public API ---------------------------------------------------------
    def parse(self, file: str | Path | BinaryIO) -> ParseResult:
        sheet_ref = self.mapping.get("sheet", 0)
        header_row = int(self.mapping.get("header_row", 1))
        wb = load_workbook(file, read_only=True, data_only=True)
        ws = wb.worksheets[sheet_ref] if isinstance(sheet_ref, int) else wb[sheet_ref]

        rows = list(ws.iter_rows(values_only=True))
        if len(rows) < header_row:
            return ParseResult()
        headers = [str(h).strip() if h is not None else "" for h in rows[header_row - 1]]
        index = {h: i for i, h in enumerate(headers)}

        result = ParseResult()
        for n, row in enumerate(rows[header_row:], start=header_row + 1):
            raw = {h: row[i] if i < len(row) else None for h, i in index.items()}
            if all(v is None or v == "" for v in raw.values()):
                continue
            try:
                tx = self.normalize_row(raw, n)
                if tx is not None:
                    result.transactions.append(tx)
            except Exception as exc:  # noqa: BLE001 - report per row, keep going
                result.errors.append(RowError(row=n, message=str(exc), raw=raw))
        wb.close()
        return result

    # --- overridable normalisation -----------------------------------------
    def normalize_row(self, raw: dict[str, Any], row_no: int) -> CanonicalTransaction | None:
        """Default mapping-driven normalisation. Override for quirks."""
        cols = self.mapping["columns"]
        src_type = self._cell(raw, cols.get("type"))
        tx_type = self.map_type(src_type)
        if tx_type is None:
            raise ValueError(f"Tipo de operación no soportado: {src_type!r}")

        return CanonicalTransaction(
            external_id=self._opt_str(self._cell(raw, cols.get("external_id"))),
            timestamp=self._parse_dt(self._cell(raw, cols.get("timestamp"))),
            type=tx_type,
            asset_in=self._sym(self._cell(raw, cols.get("asset_in"))),
            amount_in=self._dec(self._cell(raw, cols.get("amount_in"))),
            asset_out=self._sym(self._cell(raw, cols.get("asset_out"))),
            amount_out=self._dec(self._cell(raw, cols.get("amount_out"))),
            fee_asset=self._sym(self._cell(raw, cols.get("fee_asset"))),
            fee_amount=self._dec(self._cell(raw, cols.get("fee_amount"))),
            eur_value=self._dec(self._cell(raw, cols.get("eur_value"))),
            price_source=self.name,
            raw=raw,
        )

    def make_tx(
        self,
        *,
        tx_type: TransactionType,
        timestamp: datetime,
        external_id: str | None,
        eur_value: Decimal | None,
        raw: dict[str, Any],
        recv_asset: str | None = None,
        recv_amount: Decimal | None = None,
        give_asset: str | None = None,
        give_amount: Decimal | None = None,
        fee_asset: str | None = None,
        fee_amount: Decimal | None = None,
    ) -> CanonicalTransaction:
        """Assemble a canonical transaction from a received/given leg pair."""
        return CanonicalTransaction(
            external_id=external_id,
            timestamp=timestamp,
            type=tx_type,
            asset_in=recv_asset,
            amount_in=recv_amount,
            asset_out=give_asset,
            amount_out=give_amount,
            fee_asset=fee_asset,
            fee_amount=fee_amount,
            eur_value=eur_value,
            price_source=self.name,
            raw=raw,
        )

    @staticmethod
    def synth_id(*parts: Any) -> str:
        """Stable synthetic id from row content (for exports without an id)."""
        import hashlib

        joined = "|".join("" if p is None else str(p) for p in parts)
        return hashlib.sha1(joined.encode("utf-8")).hexdigest()[:20]

    def map_type(self, src_type: Any) -> TransactionType | None:
        if src_type is None:
            return None
        key = str(src_type).strip()
        type_map: dict[str, str] = self.mapping.get("type_map", {})
        mapped = type_map.get(key) or type_map.get(key.lower())
        return TransactionType(mapped) if mapped else None

    # --- cell helpers -------------------------------------------------------
    @staticmethod
    def _cell(raw: dict[str, Any], column: str | None) -> Any:
        if not column:
            return None
        return raw.get(column)

    @staticmethod
    def _opt_str(value: Any) -> str | None:
        if value is None or value == "":
            return None
        return str(value).strip()

    @staticmethod
    def _sym(value: Any) -> str | None:
        if value is None or value == "":
            return None
        return str(value).strip().upper()

    def _dec(self, value: Any) -> Decimal | None:
        if value is None or value == "":
            return None
        sep = self.mapping.get("decimal_sep", ".")
        text = str(value).strip()
        if sep == ",":
            text = text.replace(".", "").replace(",", ".")
        return abs(to_decimal(text))  # sign is implied by in/out direction

    def _parse_dt(self, value: Any) -> datetime:
        if isinstance(value, datetime):
            return value
        if value is None or value == "":
            raise ValueError("Fecha vacía")
        for fmt in self.mapping.get("date_formats", [self.mapping.get("date_format")]):
            if not fmt:
                continue
            try:
                return datetime.strptime(str(value).strip(), fmt)
            except ValueError:
                continue
        raise ValueError(f"Formato de fecha no reconocido: {value!r}")
